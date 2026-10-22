#!/usr/bin/env python3

# scripts/download_data.py

import argparse
import os
import sys
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import BadZipFile, ZipFile

import requests
from openpyxl import load_workbook

WORLD_BANK_MONTHLY_PRICES_URL = (
    "https://thedocs.worldbank.org/en/doc/"
    "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/"
    "related/CMO-Historical-Data-Monthly.xlsx"
)

SCRIPT_DIR = Path(__file__).resolve().parent
REPOSITORY_ROOT = SCRIPT_DIR.parent
DATA_ROOT = REPOSITORY_ROOT / "data" / "raw"

WORLD_BANK_OUTPUT_PATH = DATA_ROOT / "world_bank" / "CMO-Historical-Data-Monthly.xlsx"

REQUIRED_WORLD_BANK_SHEETS = {
    "Monthly Prices",
    "Monthly Indices",
    "Description",
}

XLSX_SIGNATURE = b"PK"


def download_file(
    url: str,
    destination: Path,
    *,
    force: bool = False,
    timeout: float = 60.0,
) -> Path:
    """
    Download a file atomically.

    The file is first written to a temporary location and moved into
    place only after the download succeeds.

    Args:
        url (str):
            URL of the file to download.

        destination (pathlib.Path):
            Local path where the downloaded file should be saved.

        force (bool, optional):
            If True, download the file even if it already exists.
            Defaults to False.

        timeout (float, optional):
            Timeout for the download request in seconds. Defaults to 60.0.

    Returns:
        Path:
            Path to the downloaded file.

    Raises:
        RuntimeError:
            If the temporary download path could not be created.

        requests.RequestException:
            If the download request fails.

        Exception:
            If any other error occurs during the download process.
    """

    if destination.exists() and not force:
        print(f"Already present: {destination}")
        return destination

    destination.parent.mkdir(parents=True, exist_ok=True)

    print(f"Downloading: {url}")
    print(f"Destination: {destination.resolve()}")

    temporary_path: Path | None = None

    try:
        with requests.get(
            url,
            stream=True,
            timeout=timeout,
            headers={"User-Agent": ("SyntheScale/0.1 (materials scale-up feasibility prototype)")},
        ) as response:
            response.raise_for_status()

            with NamedTemporaryFile(
                mode="wb",
                prefix=f".{destination.name}.",
                suffix=".tmp",
                dir=destination.parent,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)

                for chunk in response.iter_content(chunk_size=1024 * 1024):
                    if chunk:
                        temporary_file.write(chunk)

                temporary_file.flush()
                os.fsync(temporary_file.fileno())

        if temporary_path is None:
            raise RuntimeError("Temporary download path was not created.")

        temporary_path.replace(destination)

    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)

        raise

    print(f"Downloaded {destination.stat().st_size:,} bytes.")

    return destination


def validate_world_bank_workbook(path: Path) -> None:
    """
    Validate that the download is an expected World Bank workbook.

    Args:
        path (pathlib.Path):
            Path to the downloaded World Bank workbook.

    Raises:
        FileNotFoundError:
            If the downloaded workbook does not exist.

        ValueError:
            If the downloaded workbook is empty, not an XLSX file,
            corrupt, or missing required sheets.
    """

    if not path.exists():
        raise FileNotFoundError(f"Downloaded workbook does not exist: {path}")

    if path.stat().st_size == 0:
        raise ValueError(f"Downloaded workbook is empty: {path}")

    with path.open("rb") as handle:
        signature = handle.read(2)

    if signature != XLSX_SIGNATURE:
        raise ValueError(
            "Downloaded file is not an XLSX workbook. "
            "The server may have returned an HTML error page."
        )

    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ValueError(f"Downloaded XLSX archive is corrupt: {path}")

    except BadZipFile as error:
        raise ValueError(f"Downloaded file is not a valid XLSX archive: {path}") from error

    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        available_sheets = set(workbook.sheetnames)

    finally:
        workbook.close()

    missing_sheets = REQUIRED_WORLD_BANK_SHEETS - available_sheets

    if missing_sheets:
        missing = ", ".join(sorted(missing_sheets))

        raise ValueError(f"World Bank workbook is missing expected sheets: {missing}")

    print("Workbook validation passed.")
    print("Available sheets: " + ", ".join(sorted(available_sheets)))


def download_world_bank_data(*, force: bool = False) -> Path:
    """
    Download and validate World Bank monthly commodity prices.

    Args:
        force (bool):
            If True, re-download the dataset even if it already exists.

    Returns:
        pathlib.Path:
            Path to the downloaded and validated World Bank workbook.
    """
    path = download_file(
        WORLD_BANK_MONTHLY_PRICES_URL,
        WORLD_BANK_OUTPUT_PATH,
        force=force,
    )

    validate_world_bank_workbook(path)

    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download external datasets required by SyntheScale."
    )

    parser.add_argument(
        "--force",
        action="store_true",
        help="Replace datasets that already exist.",
    )

    parser.add_argument(
        "--world-bank",
        action="store_true",
        help="Download World Bank commodity-price data.",
    )

    return parser.parse_args()


def main() -> int:
    args = parse_args()

    download_world_bank = args.world_bank

    # If no individual dataset is selected, download all supported data.
    if not download_world_bank:
        download_world_bank = True

    try:
        if download_world_bank:
            download_world_bank_data(force=args.force)

    except (
        OSError,
        requests.RequestException,
        ValueError,
    ) as error:
        print(f"Data download failed: {error}", file=sys.stderr)
        return 1

    print("SyntheScale data setup complete.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
